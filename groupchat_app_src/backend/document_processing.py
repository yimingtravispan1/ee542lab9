"""Extract supported uploads and create hard-bounded 500-character chunks."""
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from langchain_text_splitters import CharacterTextSplitter
from pypdf import PdfReader

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 1_000_000
SUPPORTED_SUFFIXES = {'.pdf', '.txt', '.md', '.docx'}


def extract_pages(filename: str, data: bytes, *, max_file_bytes=MAX_FILE_BYTES) -> list[tuple[int, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError('Supported file types: PDF, TXT, MD, and DOCX.')
    if not data:
        raise ValueError('The uploaded file is empty.')
    if len(data) > max_file_bytes:
        raise ValueError(f'Each file must be {max_file_bytes // (1024 * 1024)} MB or smaller.')
    try:
        if suffix == '.pdf':
            if not data.lstrip().startswith(b'%PDF-'):
                raise ValueError('The file does not contain a valid PDF header.')
            reader = PdfReader(BytesIO(data))
            # Public DMV PDFs may restrict editing while allowing password-free reading.
            if reader.is_encrypted and not reader.decrypt(''):
                raise ValueError('Password-protected PDFs are not supported. Upload an unlocked copy.')
            if len(reader.pages) > 1000:
                raise ValueError('PDFs must contain no more than 1000 pages.')
            pages = [(i + 1, page.extract_text() or '') for i, page in enumerate(reader.pages)]
        elif suffix == '.docx':
            doc = DocxDocument(BytesIO(data))
            text = '\n'.join(p.text for p in doc.paragraphs)
            text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
            pages = [(1, text)]
        else:
            pages = [(1, data.decode('utf-8-sig'))]
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('Unable to read this document. Check the format; text files must use UTF-8.') from error
    pages = [(page, text.replace('\x00', '').strip()) for page, text in pages]
    total = sum(len(text) for _, text in pages)
    if total > MAX_TEXT_CHARS:
        raise ValueError('Extracted text exceeds the 1,000,000-character limit.')
    if not any(text for _, text in pages):
        raise ValueError('No readable text found. Scanned PDFs require OCR before uploading.')
    return pages


def split_pages(pages: list[tuple[int, str]]) -> list[dict]:
    # Empty separator guarantees a maximum of 500 characters, even for long
    # unbroken PDF lines. A newline-only separator does not guarantee this.
    splitter = CharacterTextSplitter(separator='', chunk_size=500, chunk_overlap=100, length_function=len, strip_whitespace=False)
    return [{'page': page, 'content': chunk} for page, text in pages for chunk in splitter.split_text(text) if chunk.strip()]
